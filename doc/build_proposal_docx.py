# -*- coding: utf-8 -*-
"""生成毕业设计开题报告初稿（docx）。"""

import os

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)),
                   "开题报告_基于简历解析与岗位匹配的智能求职辅助平台_初稿.docx")

SONG = "宋体"
HEI = "黑体"
ASCII_FONT = "Times New Roman"


# --------------------------------------------------------------------------
# 基础工具
# --------------------------------------------------------------------------
def set_run_font(run, cn=SONG, size=12, bold=False, italic=False, color=None):
    run.font.name = ASCII_FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    if color is not None:
        run.font.color.rgb = color
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), cn)
    rfonts.set(qn("w:ascii"), ASCII_FONT)
    rfonts.set(qn("w:hAnsi"), ASCII_FONT)


def setup_page(doc):
    sec = doc.sections[0]
    sec.page_width = Cm(21.0)
    sec.page_height = Cm(29.7)
    sec.top_margin = Cm(2.54)
    sec.bottom_margin = Cm(2.54)
    sec.left_margin = Cm(3.0)
    sec.right_margin = Cm(2.5)


def setup_styles(doc):
    normal = doc.styles["Normal"]
    normal.font.name = ASCII_FONT
    normal.font.size = Pt(12)
    rpr = normal.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), SONG)
    rfonts.set(qn("w:ascii"), ASCII_FONT)
    rfonts.set(qn("w:hAnsi"), ASCII_FONT)
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(0)

    for lvl, size in ((1, 16), (2, 14)):
        st = doc.styles["Heading %d" % lvl]
        st.font.name = ASCII_FONT
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = RGBColor(0, 0, 0)
        srpr = st.element.get_or_add_rPr()
        srfonts = srpr.find(qn("w:rFonts"))
        if srfonts is None:
            srfonts = OxmlElement("w:rFonts")
            srpr.append(srfonts)
        srfonts.set(qn("w:eastAsia"), HEI)
        srfonts.set(qn("w:ascii"), ASCII_FONT)
        srfonts.set(qn("w:hAnsi"), ASCII_FONT)
        pf = st.paragraph_format
        pf.first_line_indent = Pt(0)
        pf.left_indent = Pt(0)
        pf.space_before = Pt(14 if lvl == 1 else 8)
        pf.space_after = Pt(8 if lvl == 1 else 4)
        pf.line_spacing = 1.5
        pf.keep_with_next = True


def add_p(doc, text, indent=True, align=None, space_after=0, size=12,
          bold=False, cn=SONG, line_spacing=1.5):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.first_line_indent = Pt(size * 2) if indent else Pt(0)
    pf.line_spacing = line_spacing
    pf.space_before = Pt(0)
    pf.space_after = Pt(space_after)
    if align is not None:
        p.alignment = align
    r = p.add_run(text)
    set_run_font(r, cn=cn, size=size, bold=bold)
    return p


def add_h(doc, text, level):
    p = doc.add_paragraph(style="Heading %d" % level)
    r = p.add_run(text)
    set_run_font(r, cn=HEI, size=16 if level == 1 else 14, bold=True,
                 color=RGBColor(0, 0, 0))
    return p


def add_caption(doc, text):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.first_line_indent = Pt(0)
    pf.space_before = Pt(8)
    pf.space_after = Pt(3)
    pf.line_spacing = 1.0
    pf.keep_with_next = True
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    set_run_font(r, cn=SONG, size=10.5, bold=True)
    return p


def _set_cell_width(cell, cm):
    cell.width = Cm(cm)
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(cm * 567)))
    tc_w.set(qn("w:type"), "dxa")


def _shade(cell, hex_color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tc_pr.append(shd)


def remove_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement("w:%s" % edge)
        el.set(qn("w:val"), "none")
        el.set(qn("w:sz"), "0")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "auto")
        borders.append(el)
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is not None:
        layout.addprevious(borders)
    else:
        tbl_pr.append(borders)


def restart_page_numbering(section):
    sect_pr = section._sectPr
    pg_num = OxmlElement("w:pgNumType")
    pg_num.set(qn("w:start"), "1")
    sect_pr.append(pg_num)


def clear_cell_borders(cell):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement("w:%s" % edge)
        el.set(qn("w:val"), "nil")
        borders.append(el)
    tc_pr.append(borders)


def _repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def _cant_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant = OxmlElement("w:cantSplit")
    tr_pr.append(cant)


def add_table(doc, headers, rows, widths, caption=None, align_left_cols=None):
    if caption:
        add_caption(doc, caption)
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    for idx, w in enumerate(widths):
        table.columns[idx].width = Cm(w)

    hdr = table.rows[0]
    _repeat_header(hdr)
    _cant_split(hdr)
    for idx, text in enumerate(headers):
        cell = hdr.cells[idx]
        _set_cell_width(cell, widths[idx])
        _shade(cell, "F2F2F2")
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(text)
        set_run_font(r, cn=HEI, size=10.5, bold=True)

    for row_data in rows:
        row = table.add_row()
        _cant_split(row)
        for idx, text in enumerate(row_data):
            cell = row.cells[idx]
            _set_cell_width(cell, widths[idx])
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            p = cell.paragraphs[0]
            p.paragraph_format.first_line_indent = Pt(0)
            p.paragraph_format.line_spacing = 1.0
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            if align_left_cols is None or idx in align_left_cols:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(text)
            set_run_font(r, cn=SONG, size=10.5)

    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_before = Pt(0)
    spacer.paragraph_format.space_after = Pt(0)
    spacer.paragraph_format.line_spacing = 1.0
    spacer.add_run("").font.size = Pt(6)
    return table


def add_page_number(paragraph):
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_end)
    set_run_font(run, cn=SONG, size=9)


def setup_footer(doc):
    sec = doc.sections[0]
    sec.different_first_page_header_footer = True
    footer = sec.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_page_number(p)


# --------------------------------------------------------------------------
# 封面
# --------------------------------------------------------------------------
def build_cover(doc):
    for _ in range(3):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(0)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("××××大学")
    set_run_font(r, cn=HEI, size=26, bold=True)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(18)
    r = p.add_run("毕业设计（论文）开题报告")
    set_run_font(r, cn=HEI, size=30, bold=True)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(36)
    p.paragraph_format.left_indent = Cm(1.6)
    p.paragraph_format.line_spacing = 1.5
    r = p.add_run("题　目：基于简历解析与岗位匹配的")
    set_run_font(r, cn=SONG, size=15, bold=True)
    r.add_break()
    r2 = p.add_run("　　　　智能求职辅助平台的设计与实现")
    set_run_font(r2, cn=SONG, size=15, bold=True)

    for _ in range(2):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(0)

    rows = [
        ("学　　院", "____________________"),
        ("专　　业", "计算机科学与技术"),
        ("班　　级", "____________________"),
        ("学　　号", "____________________"),
        ("学生姓名", "____________________"),
        ("指导教师", "____________________"),
    ]
    table = doc.add_table(rows=len(rows), cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for i, (label, value) in enumerate(rows):
        row = table.rows[i]
        row.height = Cm(0.95)
        for j, text in enumerate((label, value)):
            cell = row.cells[j]
            _set_cell_width(cell, 3.2 if j == 0 else 7.8)
            clear_cell_borders(cell)
            p = cell.paragraphs[0]
            p.paragraph_format.line_spacing = 1.0
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT if j == 0 else WD_ALIGN_PARAGRAPH.LEFT
            r = p.add_run(("　" + text) if j == 0 else ("　" + text))
            set_run_font(r, cn=SONG, size=14)

    remove_table_borders(table)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(60)
    r = p.add_run("2026 年 11 月")
    set_run_font(r, cn=SONG, size=16)


# --------------------------------------------------------------------------
# 正文
# --------------------------------------------------------------------------
def sec_intro(doc):
    add_h(doc, "一、课题背景与研究意义", 1)
    add_h(doc, "1.1 研究背景", 2)

    add_p(doc, "近年来，高校毕业生规模持续扩大，2025 届毕业生已达 1222 万人，就业竞争日趋激烈。"
               "在招聘流程中，简历是求职者与用人单位之间的第一道关口：对企业而言，从成百上千份简历中"
               "筛选出合适人选成本高昂；对求职者而言，一份与岗位要求对齐的简历，直接决定了能否获得面试"
               "机会。然而，大量求职者仍采用“一份简历投遍所有岗位”的策略，简历内容与岗位描述（Job "
               "Description，JD）之间存在明显的信息错位，导致投递转化率低下。")

    add_p(doc, "现有的简历类工具大致可分为三类：一是模板排版类工具，解决“好看”的问题，却不解决"
               "“匹配”的问题；二是招聘平台内置的简历诊断，多基于关键词命中给出粗略提示，缺乏可解释"
               "的依据；三是通用大模型辅助写作，虽然能快速润色文字，但普遍存在三方面缺陷：其一，容易"
               "生成无法核实的量化成果，面试追问时难以自圆其说；其二，缺乏行业语境，话术与目标岗位"
               "错位；其三，缺少岗位针对性，改写结果千篇一律。")

    add_p(doc, "与此同时，相关技术条件已趋于成熟。文档智能领域出现了 LayoutLM 等兼顾文本与版面的"
               "预训练模型；中文语义向量模型（如 BGE 系列）在检索与匹配任务上表现稳定；大模型既可通过"
               "云端接口调用，也可本地部署。这使得在本科毕业设计的规模内，构建一个覆盖“简历解析—"
               "岗位匹配—诊断改写”完整闭环的系统成为可能。")

    add_p(doc, "基于上述背景，本课题拟设计并实现一个面向求职者的智能求职辅助平台：以结构化简历为"
               "数据底座，把非结构化的简历文档转化为可计算的能力画像；以岗位描述为参照，输出可解释的"
               "匹配评分、能力差距与风险提示；以证据约束的方式提供改写建议，避免大模型凭空编造；并通过"
               "离线评测验证方法的有效性。")

    add_h(doc, "1.2 研究意义", 2)
    add_p(doc, "本课题的意义主要体现在以下三个方面。")
    add_p(doc, "（1）理论意义。把简历从“一份文档”抽象为“结构化能力画像”，为人岗匹配提供了可计算、"
               "可解释、可评测的中间表示；同时探索证据约束下的受控文本生成方法，为缓解简历改写中的"
               "事实性幻觉提供一种工程化思路。")
    add_p(doc, "（2）实践意义。系统能够为求职者输出可操作的差距清单与修改建议，降低其获取“HR 视角”"
               "反馈的成本；对高校就业指导部门而言，群体层面的能力画像与岗位匹配数据，可为课程设置与"
               "就业辅导提供参考。")
    add_p(doc, "（3）应用价值。系统采用前后端分离与容器化部署，既可作为独立的求职辅助工具运行，也可"
               "作为文档智能与语义匹配的教学案例，具备一定的可复用性。")


def sec_related(doc):
    add_h(doc, "二、国内外研究现状", 1)

    add_h(doc, "2.1 简历解析", 2)
    add_p(doc, "简历解析的目标是把非结构化的简历文档转换为结构化字段，包括基本信息、教育背景、"
               "工作与实习经历、项目经历、专业技能、获奖情况等。现有方法大致可分为四类。")
    add_p(doc, "（1）规则与模板方法。依据固定版式编写正则表达式与关键词词典，直接定位字段。该方法"
               "可控、可解释、无需标注数据，但对版式变化极为敏感，维护成本随模板数量增长而快速上升。")
    add_p(doc, "（2）序列标注方法。把字段抽取建模为命名实体识别任务，采用 BiLSTM-CRF、BERT 等模型。"
               "该类方法泛化能力较强，但依赖一定规模的标注数据，且难以利用版式信息。")
    add_p(doc, "（3）文档图像理解方法。LayoutLM、LayoutLMv3 等工作将文本、位置与图像统一预训练，"
               "在表单、票据等文档上效果显著。此类模型能力较强，但参数量大、推理成本高，针对中文简历"
               "的领域适配仍不充分。")
    add_p(doc, "（4）大模型抽取方法。借助大模型的零样本或少样本能力直接输出结构化结果，开发成本低，"
               "但输出格式不稳定、存在字段幻觉，且难以给出可靠的置信度。")
    add_p(doc, "总体来看，单一方法难以同时兼顾鲁棒性与可控性，工程实践中更常见的是“规则优先、模型"
               "兜底、人工校正”的混合路线。中文简历版式多样（双栏、表格、图标、扫描件），公开可用的"
               "中文解析工具与字段级评测数据均相对缺乏，这为本课题留出了改进空间。")

    add_h(doc, "2.2 人岗匹配", 2)
    add_p(doc, "（1）关键词与统计检索方法。以 TF-IDF、BM25 为代表，通过词项匹配计算简历与岗位描述的"
               "相关度。该类方法计算高效、结果可解释，但无法处理同义表达与语义相似，例如“前端开发”"
               "与“Web 前端”、“JS”与“JavaScript”。")
    add_p(doc, "（2）语义向量方法。从 word2vec 到 BERT，再到 Sentence-BERT 与 BGE 系列中文向量模型，"
               "文本表示能力不断提升。当前主流做法是“双塔召回 + 交叉编码重排”，在语义泛化上优势明显，"
               "但可解释性较弱，且存在领域漂移问题。")
    add_p(doc, "（3）知识图谱方法。通过构建技能本体及上下位关系，支持“掌握 Spring Boot 是否满足"
               "Spring Cloud 要求”这类推理。该类方法可解释性好，但本体构建与维护成本高。")
    add_p(doc, "（4）融合与重排方法。将多路召回结果通过倒数排名融合（RRF）等策略合并，再由重排模型"
               "精排，是工业界的主流方案。")
    add_p(doc, "现有研究多面向招聘平台的海量候选场景，强调召回效率；而面向单个求职者的“简历—岗位”"
               "诊断场景，更强调解释性与可操作性，二者的优化目标并不完全相同。此外，公开的中文简历与"
               "岗位配对标注数据稀缺，使离线评测难以复现，这也是现有商业产品普遍只给分数、不给依据的"
               "原因之一。")

    add_h(doc, "2.3 简历优化与文本生成", 2)
    add_p(doc, "早期简历优化以规则检查为主，例如关键词覆盖率、STAR 结构完整性、量化成果缺失等。"
               "大模型出现后，改写与扩写成为可能，但事实性幻觉成为主要风险。检索增强生成（RAG）与"
               "引用溯源是常见的缓解手段，其核心思想是让生成内容有据可查。然而在简历场景中，“依据”"
               "往往存在于求职者本人而非外部知识库，因此需要把生成内容与用户已确认的事实字段绑定，"
               "这正是本课题设计“证据链 + 人机确认”机制的出发点。")

    add_h(doc, "2.4 现有系统与产品", 2)
    add_p(doc, "从产品形态看，现有简历相关系统可分为四类，其能力与不足如表 2-1 所示。")
    add_table(
        doc,
        ["类型", "代表产品", "主要能力", "不足"],
        [
            ["模板排版类", "超级简历、五百丁等", "模板丰富、排版美观、导出方便",
             "只解决呈现问题，不做岗位匹配"],
            ["招聘平台类", "BOSS 直聘、智联招聘等", "投递、沟通、企业筛选一体化",
             "诊断面向企业侧，颗粒度粗，求职者缺少可操作建议"],
            ["关键词诊断类", "Jobscan、Rezi 等", "按 JD 计算关键词覆盖率",
             "以英文语境为主，规则不透明，难以说明具体差距"],
            ["AI 简历类", "小坏简历等", "JD 诊断、行业化扩写、证据确认",
             "评分依赖大模型，缺少离线评测与可复现机制，系统与商业功能耦合"],
        ],
        [2.8, 4.0, 4.6, 4.1],
        caption="表 2-1　现有简历相关系统对比",
        align_left_cols={0, 1, 2, 3},
    )
    add_p(doc, "可以看出，现有产品在交互体验与功能完整度上已较为成熟，但在可解释性、可评测性与"
               "可复现性方面普遍不足。")

    add_h(doc, "2.5 小结", 2)
    add_p(doc, "综上，本课题不以“再造一个简历模板站”为目标，而是围绕“结构化表示—可解释匹配—"
               "证据约束改写—离线评测”这条主线展开，力求在有限的工作量内形成可验证的技术贡献。")


def sec_content(doc):
    add_h(doc, "三、研究内容与研究目标", 1)

    add_h(doc, "3.1 研究内容", 2)
    add_p(doc, "本课题围绕“简历—岗位”这一对核心对象，开展以下五方面工作。")
    add_p(doc, "（1）结构化简历模型与多版式解析。设计可扩展的简历 JSON Schema，覆盖基本信息、教育"
               "背景、工作经历、项目经历、专业技能、获奖情况等区块；实现 PDF 与 DOCX 文档的文本及"
               "坐标抽取、版面分块、阅读顺序重排、段落分类与字段抽取；引入置信度机制与人工校正闭环，"
               "把用户确认结果回流为增量标注数据。")
    add_p(doc, "（2）岗位描述解析与技能标准化。从 JD 中抽取硬性要求（学历、工作年限、技能）与"
               "软性要求，区分“必须”与“优先”；构建技能词典及其同义词、上下位关系，实现技能归一与"
               "熟练度分级。")
    add_p(doc, "（3）可解释人岗匹配与差距诊断。融合技能覆盖、语义相似、经历相关与教育背景四类信号，"
               "形成可拆解的评分；针对每一项得分输出证据，包括命中的简历片段、对应的岗位要求句，以及"
               "缺失项与风险点。")
    add_p(doc, "（4）证据约束的简历改写与人机确认。基于目标岗位与差距项生成改写候选，以句级差异"
               "（diff）方式呈现；对新增的量化数据或职责描述标注来源，无来源的内容标记为“待确认”，"
               "必须经用户补充或确认后方可写入。")
    add_p(doc, "（5）系统实现与评测。完成平台的设计、开发与部署；构建中文简历与岗位配对的小规模"
               "评测集，从字段抽取、匹配一致性、建议质量三个维度开展离线评测与消融实验。")

    add_h(doc, "3.2 研究目标", 2)
    add_p(doc, "本课题的总体目标是交付一套可部署、可演示、可评测的智能求职辅助平台。具体量化指标"
               "如表 3-1 所示。")
    add_table(
        doc,
        ["指标", "目标值", "验证方式"],
        [
            ["支持文档格式", "PDF、DOCX", "功能测试"],
            ["字段抽取 F1 值", "≥ 0.85", "自建 100 份测试集"],
            ["匹配分与人工评分一致性", "Spearman ≥ 0.70", "50 组人工标注对照"],
            ["单份简历解析时延", "≤ 10 s", "性能测试"],
            ["单次匹配计算时延", "≤ 2 s（不含生成）", "性能测试"],
            ["改写建议事实错误率", "≤ 5%", "人工核查 100 条建议"],
        ],
        [5.5, 4.5, 5.5],
        caption="表 3-1　研究目标与验证方式",
        align_left_cols={0, 1, 2},
    )
    add_p(doc, "上述指标均在自建数据集与单机部署环境下测量，不作为生产级性能承诺。")

    add_h(doc, "3.3 拟解决的关键问题", 2)
    add_p(doc, "（1）版式鲁棒性问题。双栏排版、表格嵌套、图标与时间轴等元素会导致文本提取后的阅读"
               "顺序错乱，进而使字段归属错误。需要在缺乏版面标注数据的条件下，设计轻量而有效的分块与"
               "重排策略。")
    add_p(doc, "（2）技能归一问题。同一技能存在中英文混写、缩写、别名与层级差异，且求职者对熟练度的"
               "表述高度自由，需要建立可维护的词典与归一流程。")
    add_p(doc, "（3）可解释与可评测问题。匹配分数必须能够拆解到具体证据，且算法版本、权重与输入快照"
               "需可追溯，否则无法复现，也无法支撑论文中的对比实验。")
    add_p(doc, "（4）生成可控问题。大模型在改写时倾向于补全细节，容易引入未经证实的数据，需要从"
               "提示词约束、结构化输入与后置校验三个层面共同抑制。")


def sec_requirement(doc):
    add_h(doc, "四、系统需求分析", 1)

    add_h(doc, "4.1 用户角色", 2)
    add_p(doc, "系统面向校内求职场景设计，主要角色及其职责如表 4-1 所示。企业端仅保留轻量的脱敏"
               "查看能力，不作为本期建设重点。")
    add_table(
        doc,
        ["角色", "主要职责", "典型场景"],
        [
            ["求职者（学生）", "维护个人档案与简历，上传解析，选择岗位，查看诊断，采纳建议，导出与投递记录",
             "秋招前针对目标岗位修改简历"],
            ["管理员", "用户与权限管理，岗位库维护，技能词典维护，模板管理，日志审计，评测集管理",
             "批量导入岗位、处理异常账号"],
            ["企业用户（可选）", "脱敏查看投递简历，记录查看日志",
             "校园招聘初筛"],
        ],
        [3.2, 6.3, 6.0],
        caption="表 4-1　用户角色与职责",
        align_left_cols={0, 1, 2},
    )

    add_h(doc, "4.2 功能需求", 2)
    add_p(doc, "系统功能划分为简历、岗位、诊断、改写、投递与系统管理六个部分，核心功能点如表 4-2 "
               "所示。")
    add_table(
        doc,
        ["编号", "模块", "功能说明", "优先级"],
        [
            ["F01", "账号", "注册、登录、个人信息与求职意向维护", "高"],
            ["F02", "简历", "分区块表单编辑，实时预览，多版本管理", "高"],
            ["F03", "简历", "上传 PDF/DOCX，异步解析为结构化字段，低置信字段高亮校正", "高"],
            ["F04", "简历", "服务端渲染导出 PDF，保证预览与导出的一致性", "高"],
            ["F05", "岗位", "岗位录入与 JD 解析，技能标签提取，岗位检索与收藏", "高"],
            ["F06", "诊断", "计算总分与分项得分，输出命中、缺失、风险与证据", "高"],
            ["F07", "改写", "针对差距项生成改写候选，句级 diff，逐条采纳或拒绝", "高"],
            ["F08", "改写", "关键信息缺失时生成澄清问题，引导用户补充", "中"],
            ["F09", "投递", "投递记录与状态看板（已投、笔试、面试、结果）", "中"],
            ["F10", "管理", "用户、岗位、技能词典、模板、操作日志与评测集管理", "中"],
        ],
        [1.6, 2.6, 8.3, 3.0],
        caption="表 4-2　系统功能需求清单",
        align_left_cols={0, 1, 2, 3},
    )

    add_h(doc, "4.3 非功能需求", 2)
    add_p(doc, "（1）性能。单份简历解析在 10 s 内完成，匹配计算在 2 s 内返回；解析任务采用异步"
               "处理，避免阻塞前端。")
    add_p(doc, "（2）安全性。遵循个人信息最小化原则，简历原件加密存储，导出文件可加水印；企业侧"
               "查看需脱敏并记录访问日志；接口实施基于角色的访问控制。")
    add_p(doc, "（3）可靠性。解析失败时给出明确提示并支持转手工录入；关键操作具备幂等性；任务失败"
               "可重试。")
    add_p(doc, "（4）易用性。核心流程（上传—选岗—诊断—改写—导出）不超过五步；关键结论均有文字"
               "解释，不依赖颜色或图形传达信息。")
    add_p(doc, "（5）可维护性。简历 Schema、技能词典、评分权重均以配置方式管理，支持在不修改代码的"
               "前提下调整。")

    add_h(doc, "4.4 核心用例", 2)
    add_p(doc, "用例一：上传简历并生成诊断报告。求职者登录后上传 PDF 或 DOCX 简历，系统异步完成"
               "解析并返回结构化结果；对置信度较低的字段以高亮方式提示确认；用户选择目标岗位后，系统"
               "计算匹配评分并生成报告，报告包含总分、分项得分、命中技能、缺失技能、经历匹配情况与"
               "风险提示。异常流程包括：文件格式不支持、文件超过大小限制、解析失败（提示转为手工录入）、"
               "岗位库中无对应岗位（支持粘贴 JD 进行临时分析）。")
    add_p(doc, "用例二：针对岗位改写一段项目经历。用户在诊断报告中选择“项目描述过于笼统”这一建议，"
               "进入改写界面；系统基于该段经历的结构化字段与岗位要求生成候选，以句级 diff 展示新增、"
               "删除与修改；涉及新增量化数据的句子被标记为“待确认”，需用户补充具体数值后方可采纳；"
               "用户逐条采纳或拒绝，采纳后生成新的简历版本，旧版本保留可回溯。")


def sec_design(doc):
    add_h(doc, "五、系统总体设计", 1)

    add_h(doc, "5.1 总体架构", 2)
    add_p(doc, "系统采用“前后端分离 + 算法服务独立部署”的分层结构。业务应用与算法服务通过 REST "
               "接口和消息队列解耦：解析任务耗时较长，采用异步任务加状态轮询；匹配与生成采用同步接口，"
               "生成过程支持流式返回。总体架构如表 5-1 所示。")
    add_table(
        doc,
        ["层次", "主要组成", "职责"],
        [
            ["表现层", "Vue3 + TypeScript + Element Plus", "简历编辑、报告可视化、diff 交互、投递看板"],
            ["应用层", "Spring Boot 3 + MyBatis-Plus + Sa-Token", "用户与权限、简历与岗位管理、任务编排、报告落库"],
            ["算法层", "Python + FastAPI", "文档解析、技能归一、匹配打分、改写生成"],
            ["数据层", "MySQL、Redis、MinIO、pgvector", "结构化数据、缓存、简历原件、向量检索"],
            ["基础设施", "Docker Compose、Nginx", "一键部署、反向代理、静态资源托管"],
        ],
        [2.8, 5.4, 7.3],
        caption="表 5-1　系统总体架构",
        align_left_cols={0, 1, 2},
    )
    add_p(doc, "该结构的考虑是：业务逻辑以 Java 实现，便于权限、事务与后台管理的工程化；解析与"
               "算法部分依赖 Python 生态，独立部署可避免在 JVM 中重复造轮子，也便于单独调试与替换模型。")

    add_h(doc, "5.2 功能模块划分", 2)
    add_p(doc, "按业务内聚性划分为七个模块，如表 5-2 所示。")
    add_table(
        doc,
        ["模块", "包含子功能", "说明"],
        [
            ["简历中心", "在线编辑、上传解析、版本管理、导出", "系统的数据入口与核心资产"],
            ["岗位中心", "岗位录入、JD 解析、技能标签、检索收藏", "匹配的参照系"],
            ["诊断中心", "匹配计算、差距分析、风险提示、报告导出", "系统的核心价值输出"],
            ["改写中心", "候选生成、diff 对比、证据确认、追问", "从诊断到行动的闭环"],
            ["投递中心", "投递记录、状态流转、提醒", "求职过程管理"],
            ["系统管理", "用户、角色、词典、模板、日志", "后台支撑"],
            ["评测中心", "评测集管理、批量评测、指标统计", "支撑论文实验，可仅对管理员开放"],
        ],
        [3.2, 6.3, 6.0],
        caption="表 5-2　系统功能模块划分",
        align_left_cols={0, 1, 2},
    )

    add_h(doc, "5.3 数据模型设计", 2)
    add_p(doc, "系统核心实体及其关键字段如表 5-3 所示。")
    add_table(
        doc,
        ["实体", "关键字段", "说明"],
        [
            ["user", "id、账号、密码摘要、角色、求职意向", "用户主表"],
            ["resume", "id、user_id、标题、当前版本、状态", "一份简历"],
            ["resume_version", "id、resume_id、版本号、schema_json、创建时间", "版本快照，支持回溯"],
            ["resume_block", "id、version_id、区块类型、原文、起止偏移、置信度", "段落级冗余，支撑证据定位"],
            ["job", "id、标题、公司、JD 原文、解析结果、状态", "岗位主表"],
            ["skill", "id、标准名、类别、父技能", "技能词典与上下位"],
            ["skill_alias", "id、skill_id、别名、来源", "同义词与缩写"],
            ["match_report", "id、resume_version、job_id、总分、分项、证据、算法版本", "报告快照，保证可复现"],
            ["suggestion", "id、report_id、段落、原文、候选、证据、状态", "改写建议与采纳状态"],
            ["application", "id、user_id、job_id、状态、时间线", "投递记录"],
        ],
        [3.4, 7.0, 5.1],
        caption="表 5-3　核心数据实体",
        align_left_cols={0, 1, 2},
    )
    add_p(doc, "设计上有两点考虑：其一，简历正文以 JSON 存储以便整体读写，同时将段落冗余到"
               "resume_block 表，用字符偏移支撑“建议—原文”的精确对应；其二，匹配报告保存算法版本、"
               "权重配置与输入快照，使任一历史报告都可复算，这是后续评测与对比实验的基础。")

    add_h(doc, "5.4 关键流程设计", 2)
    add_p(doc, "（1）解析流程。上传文件 → 格式与大小校验、去重 → 转存 MinIO → 创建解析任务 → "
               "抽取文本与坐标 → 版面分块与阅读顺序重排 → 段落分类 → 字段抽取与置信度计算 → 结构化"
               "JSON 落库 → 低置信字段标记 → 前端校正 → 版本固化。")
    add_p(doc, "（2）匹配流程。加载简历版本与岗位 → 技能抽取与归一 → 技能覆盖计算 → 段落向量召回"
               "与重排 → 经历相关度计算 → 加权汇总 → 证据组装 → 报告快照落库。")
    add_p(doc, "（3）改写流程。选择目标段落与差距项 → 组装提示词（结构化字段、岗位要求、行业词表、"
               "禁止事项）→ 生成候选 → 句级对齐与 diff → 证据绑定与“待确认”标记 → 用户确认 → "
               "写入新版本。")


def sec_keytech(doc):
    add_h(doc, "六、关键技术路线", 1)

    add_h(doc, "6.1 简历解析", 2)
    add_p(doc, "解析是系统的第一道工序，直接决定后续匹配的质量。拟采用“规则优先、模型兜底、人工"
               "校正”的混合路线，具体步骤如下。")
    add_p(doc, "（1）文本与坐标抽取。PDF 使用 PyMuPDF 逐页抽取文本片段，记录包围盒、字号与字体；"
               "DOCX 使用 python-docx 解析段落与表格结构，保留层级信息。")
    add_p(doc, "（2）版面分块。依据文本片段的横坐标分布做聚类，判断单栏或双栏；依据字号与加粗识别"
               "标题；对表格按行列还原为二维结构，避免单元格文本被拉平。")
    add_p(doc, "（3）阅读顺序重排。栏内按纵坐标排序，栏间按横坐标排序，跨栏标题单独处理，形成符合"
               "人类阅读习惯的文本序列。")
    add_p(doc, "（4）段落分类。先以关键词表匹配区块标题（教育背景、工作经历、项目经历、专业技能、"
               "自我评价等），未命中时用轻量文本分类模型兜底，将段落归入对应区块。")
    add_p(doc, "（5）字段抽取。时间区间使用正则匹配（如 2023.07—2024.09、2023 年 7 月至今）；"
               "学校、公司、专业使用词典匹配；技能使用技能词典与模糊匹配；学历、绩点、排名等使用"
               "规则模板。仅在规则与模型均无法确定时调用大模型兜底，并要求其输出字段级置信度。")
    add_p(doc, "（6）置信度与校正。规则精确命中记为高置信，分类模型取概率值，大模型结果记为中等"
               "置信；低于阈值的字段在前端高亮，用户确认后回写，并作为增量标注数据留存。")

    add_h(doc, "6.2 技能标准化", 2)
    add_p(doc, "技能标准化是匹配的前提，主要包括三个环节。")
    add_p(doc, "（1）技能词典构建。以公开的技能分类为骨架，整理名称、类别与上下位关系，例如"
               "“Spring Cloud”隶属于“微服务框架”，“Vue”隶属于“前端框架”。")
    add_p(doc, "（2）别名归一。按“精确匹配 → 别名匹配 → 编辑距离或向量相似 → 人工确认”的顺序"
               "逐级降级，覆盖中英文混写、缩写与常见误写。")
    add_p(doc, "（3）熟练度抽取。从“了解、熟悉、熟练、精通”等表述以及项目角色、使用年限中推断"
               "熟练度，映射为入门、熟练、精通三级，用于技能覆盖的加权计算。")
    add_p(doc, "岗位侧同样执行上述流程，并额外区分“必须”与“优先”：含有“必须、要求、至少”等"
               "表述的记为硬性要求，含有“优先、加分、了解者优先”的记为加分项。")

    add_h(doc, "6.3 匹配算法", 2)
    add_p(doc, "匹配采用多信号加权模型，兼顾可解释性与语义泛化能力。技能分定义为需求技能的加权"
               "覆盖率：")
    add_p(doc, "S_skill = Σ(w_i × h_i) / Σ(w_i)", indent=False,
          align=WD_ALIGN_PARAGRAPH.CENTER)
    add_p(doc, "其中 w_i 为第 i 项需求技能的权重，取该技能在岗位库中的逆文档频率；h_i 为命中系数，"
               "完全命中取 1，仅命中下位技能取 0.6，缺失取 0。语义分 S_sem 由简历段落与岗位要求段落"
               "的向量余弦相似度计算，取前 k 个段落的均值；召回阶段使用 BM25 与向量检索双路召回，"
               "并以倒数排名融合（RRF）合并结果。经历分 S_exp 综合时间重叠度、行业与岗位相似度、"
               "职责动词匹配度计算。教育分 S_edu 依据学历层次与专业相关度给分。")
    add_p(doc, "总分按如下加权方式合成：")
    add_p(doc, "S = 0.40 × S_skill + 0.25 × S_sem + 0.20 × S_exp + 0.15 × S_edu",
          indent=False, align=WD_ALIGN_PARAGRAPH.CENTER)
    add_p(doc, "权重以配置项形式管理，并随报告一同存为快照。相比直接由大模型给出一个总分，该方案"
               "的优势在于：每一分都能追溯到具体证据，调整权重后可立即观察分数变化，且同一输入在"
               "固定版本下结果可复现，便于开展对比实验。")

    add_h(doc, "6.4 证据约束的改写", 2)
    add_p(doc, "改写环节的核心目标是“提升表达质量，但不引入虚假事实”，拟从三个层面约束。")
    add_p(doc, "（1）输入约束。提示词仅包含结构化字段、目标岗位要求与行业词表，不传入与目标段落"
               "无关的内容，降低模型自由发挥的空间。")
    add_p(doc, "（2）生成约束。明确禁止新增数字、比例、排名等量化信息，除非用户已提供；要求输出"
               "逐句结果并标注每句依据的字段编号。")
    add_p(doc, "（3）后置校验。对生成结果做规则检查：出现未在原文与用户输入中出现过的数字即标记为"
               "“待确认”；与原文事实冲突的表述直接拒绝；通过校验的内容才允许写入简历。")
    add_p(doc, "差异展示采用句级对齐：先按标点与换行切分句子，再用序列匹配算法对齐原文与候选，"
               "分别标记新增、删除与修改，用户可逐句采纳。采纳结果写入新版本，旧版本保留，形成可"
               "回溯的修改历史。")

    add_h(doc, "6.5 评测方法", 2)
    add_p(doc, "为避免“自说自话”，本课题设计独立的离线评测方案。")
    add_p(doc, "（1）数据集。以公开简历数据集为基础进行中文化适配，并自建 50 至 100 组中文简历与"
               "岗位配对样本，由人工标注匹配等级（1 至 5 分）与字段真值。")
    add_p(doc, "（2）指标。解析任务使用字段级精确率、召回率与 F1 值；匹配任务使用 Spearman 秩相关"
               "系数衡量算法分与人工分的一致性；改写任务统计建议采纳率，并由人工对相关性、真实性、"
               "可读性三个维度盲评。")
    add_p(doc, "（3）消融实验。分别移除语义分、移除规则技能分、移除大模型兜底，观察各项指标的"
               "变化，以说明各组件的贡献。")
    add_p(doc, "（4）可复现性。固定随机种子，记录模型版本、提示词版本与权重配置，保证实验可重复。")

    add_h(doc, "6.6 技术选型", 2)
    add_p(doc, "系统技术选型如表 6-1 所示，选型原则是成熟稳定、文档完善、便于单机部署。")
    add_table(
        doc,
        ["层次", "技术方案", "选择理由"],
        [
            ["前端", "Vue3 + TypeScript + Vite + Element Plus", "生态成熟，组件丰富，适合表单密集型界面"],
            ["业务后端", "Spring Boot 3 + MyBatis-Plus + Sa-Token", "工程化能力强，权限与事务处理方便"],
            ["算法服务", "Python + FastAPI + PyMuPDF + sentence-transformers", "文档与向量生态完善，接口轻量"],
            ["数据库", "MySQL 8", "结构化数据与事务"],
            ["缓存与队列", "Redis + RabbitMQ", "会话缓存、异步解析任务"],
            ["对象存储", "MinIO", "简历原件与导出文件"],
            ["向量检索", "pgvector", "与 MySQL 之外的轻量方案，避免引入独立向量库"],
            ["部署", "Docker Compose + Nginx", "一键启动，便于答辩演示与迁移"],
            ["大模型", "OpenAI 兼容接口 / 本地 Ollama", "可切换，避免绑定单一供应商"],
        ],
        [2.4, 7.1, 6.0],
        caption="表 6-1　技术选型",
        align_left_cols={0, 1, 2},
    )


def sec_innovation(doc):
    add_h(doc, "七、创新点与特色", 1)
    add_p(doc, "（1）结构化表示与证据定位相结合的可解释匹配。区别于直接输出一个匹配分数的做法，"
               "本系统将分数拆解为技能、语义、经历、教育四个分项，每个分项均可定位到简历原文与岗位"
               "要求句，使“为什么是这个分数”可被追问、可被验证。")
    add_p(doc, "（2）置信度驱动的人机协同解析。对低置信字段主动请求用户确认，而非静默采用模型"
               "结果；用户校正数据可回流为增量标注，使解析能力随使用逐步提升。")
    add_p(doc, "（3）证据约束的改写机制。通过输入约束、生成约束与后置校验三层控制，配合句级 diff "
               "与“待确认”标记，在保留大模型表达能力的同时抑制事实性幻觉，这一点与现有产品的"
               "“一键润色”形成明显差异。")
    add_p(doc, "（4）面向中文简历的离线评测方案。构建小规模人工标注评测集，给出字段抽取、匹配一致性"
               "与建议质量的量化结果，并进行消融实验，弥补现有产品“只给结论、不给依据”的不足。")
    add_p(doc, "（5）工程层面的可复现设计。匹配报告保存算法版本、权重与输入快照，任一历史结论均可"
               "复算，为论文实验与系统迭代提供支撑。")


def sec_schedule(doc):
    add_h(doc, "八、进度安排", 1)
    add_p(doc, "本课题自 2026 年 10 月启动，计划于 2027 年 5 月完成答辩，具体安排如表 8-1 所示。")
    add_table(
        doc,
        ["时间", "阶段", "主要任务", "阶段产出"],
        [
            ["2026.10—2026.11", "选题与开题", "文献调研、需求分析、技术验证", "开题报告、原型"],
            ["2026.11—2026.12", "数据与解析", "数据准备、解析模块、简历编辑与导出", "解析服务、简历中心"],
            ["2026.12—2027.01", "岗位与匹配", "JD 解析、技能词典、匹配算法", "岗位中心、诊断报告"],
            ["2027.01—2027.02", "改写与联调", "改写模块、diff 交互、系统集成", "完整业务闭环"],
            ["2027.02—2027.03", "测试与中期", "功能测试、性能优化、中期检查", "测试报告、中期材料"],
            ["2027.03—2027.04", "评测与论文", "离线评测、消融实验、论文撰写", "实验数据、论文初稿"],
            ["2027.04—2027.05", "修改与答辩", "查重、修改、答辩准备", "定稿、答辩材料"],
        ],
        [3.4, 3.0, 5.6, 3.5],
        caption="表 8-1　进度安排",
        align_left_cols={0, 1, 2, 3},
    )
    add_p(doc, "进度安排中预留约两周缓冲，用于应对解析效果不达预期、模型接口不稳定等风险。其中"
               "解析模块是整条链路的基础，计划在 2026 年 12 月底前完成最小可用版本，避免后期返工。")


def sec_result(doc):
    add_h(doc, "九、预期成果", 1)
    add_p(doc, "（1）可部署的智能求职辅助平台一套，包含源代码、数据库脚本、接口文档与部署说明。")
    add_p(doc, "（2）毕业设计（论文）一篇，围绕结构化解析、可解释匹配与证据约束改写展开论述。")
    add_p(doc, "（3）中文简历与岗位配对评测集一份，规模 50 至 100 组，附标注规范。")
    add_p(doc, "（4）视进度情况，整理软件著作权申请材料一套。")


def sec_reference(doc):
    add_h(doc, "十、参考文献", 1)
    refs = [
        "[1] Devlin J, Chang M W, Lee K, et al. BERT: Pre-training of deep bidirectional transformers "
        "for language understanding[C]//Proceedings of the 2019 Conference of the North American "
        "Chapter of the Association for Computational Linguistics. Minneapolis: ACL, 2019: 4171-4186.",
        "[2] Huang Y, Lv T, Cui L, et al. LayoutLMv3: Pre-training for document AI with unified text "
        "and image masking[C]//Proceedings of the 30th ACM International Conference on Multimedia. "
        "New York: ACM, 2022: 4083-4091.",
        "[3] Reimers N, Gurevych I. Sentence-BERT: Sentence embeddings using Siamese BERT-networks[C]"
        "//Proceedings of the 2019 Conference on Empirical Methods in Natural Language Processing. "
        "Hong Kong: ACL, 2019: 3982-3992.",
        "[4] Chen J, Xiao S, Zhang P, et al. BGE M3-Embedding: Multi-lingual, multi-functionality, "
        "multi-granularity text embeddings through self-knowledge distillation[J]. arXiv preprint "
        "arXiv:2402.03216, 2024.",
        "[5] Robertson S, Zaragoza H. The probabilistic relevance framework: BM25 and beyond[J]. "
        "Foundations and Trends in Information Retrieval, 2009, 3(4): 333-389.",
        "[6] Cormack G V, Clarke C L A, Buettcher S. Reciprocal rank fusion outperforms Condorcet and "
        "individual rank learning methods[C]//Proceedings of the 32nd International ACM SIGIR "
        "Conference on Research and Development in Information Retrieval. New York: ACM, 2009: 758-759.",
        "[7] Lewis P, Perez E, Piktus A, et al. Retrieval-augmented generation for knowledge-intensive "
        "NLP tasks[C]//Advances in Neural Information Processing Systems 33. 2020: 9459-9474.",
        "[8] Vaswani A, Shazeer N, Parmar N, et al. Attention is all you need[C]//Advances in Neural "
        "Information Processing Systems 30. 2017: 5998-6008.",
        "[9] 李航. 统计学习方法[M]. 2 版. 北京: 清华大学出版社, 2019.",
        "[10] 周志华. 机器学习[M]. 北京: 清华大学出版社, 2016.",
        "[11] 全国信息安全标准化技术委员会. 信息安全技术 个人信息安全规范: GB/T 35273—2020[S]. "
        "北京: 中国标准出版社, 2020.",
        "[12] 中华人民共和国个人信息保护法[Z]. 2021-08-20.",
    ]
    for ref in refs:
        p = doc.add_paragraph()
        pf = p.paragraph_format
        pf.first_line_indent = Cm(-0.85)
        pf.left_indent = Cm(0.85)
        pf.line_spacing = 1.5
        pf.space_before = Pt(0)
        pf.space_after = Pt(0)
        r = p.add_run(ref)
        set_run_font(r, cn=SONG, size=10.5)


def build():
    doc = Document()
    setup_page(doc)
    setup_styles(doc)
    setup_footer(doc)
    doc.core_properties.title = "基于简历解析与岗位匹配的智能求职辅助平台的设计与实现"
    doc.core_properties.subject = "毕业设计（论文）开题报告"
    doc.core_properties.author = ""

    build_cover(doc)
    body_section = doc.add_section(WD_SECTION.NEW_PAGE)
    body_section.different_first_page_header_footer = False
    restart_page_numbering(body_section)

    sec_intro(doc)
    sec_related(doc)
    sec_content(doc)
    sec_requirement(doc)
    sec_design(doc)
    sec_keytech(doc)
    sec_innovation(doc)
    sec_schedule(doc)
    sec_result(doc)
    sec_reference(doc)

    doc.save(OUT)
    return OUT


if __name__ == "__main__":
    path = build()
    # 简单校验
    d = Document(path)
    h1 = [p.text for p in d.paragraphs if p.style.name == "Heading 1"]
    h2 = [p.text for p in d.paragraphs if p.style.name == "Heading 2"]
    print("saved:", path)
    print("size:", os.path.getsize(path))
    print("paragraphs:", len(d.paragraphs), "tables:", len(d.tables))
    print("H1:", len(h1), h1)
    print("H2:", len(h2))
    chars = sum(len(p.text) for p in d.paragraphs)
    for t in d.tables:
        for row in t.rows:
            for c in row.cells:
                chars += len(c.text)
    print("approx chars:", chars)
