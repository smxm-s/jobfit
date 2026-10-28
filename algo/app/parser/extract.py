"""文本与坐标抽取：PDF（PyMuPDF）与 DOCX（python-docx）。"""

import io
from dataclasses import dataclass, field
from typing import List, Optional, Tuple


class ParseBusinessError(Exception):
    """解析业务异常，由 main.py 转成对应 HTTP 状态与错误码。"""

    def __init__(self, code: str, message: str, http_status: int = 422) -> None:
        super().__init__(message)
        self.code = code
        self.message = message
        self.http_status = http_status


@dataclass
class Span:
    text: str
    bbox: Tuple[float, float, float, float]
    size: float
    font: str
    bold: bool
    page_no: int
    block_no: int
    style: str = ""


@dataclass
class ExtractedDocument:
    spans: List[Span] = field(default_factory=list)
    tables: List[List[List[str]]] = field(default_factory=list)
    text: str = ""
    page_count: int = 0
    doc_type: str = "PDF"
    page_width: float = 0.0


def _is_bold(font: str, flags: int) -> bool:
    lowered = (font or "").lower()
    return bool(flags & 16) or any(
        token in lowered for token in ("bold", "semibold", "heavy", "black", "hei", "黑体")
    )


def _pdf_line_text(spans: List[Span]) -> str:
    ordered = sorted(spans, key=lambda item: item.bbox[0])
    parts: List[str] = []
    prev: Optional[Span] = None
    for span in ordered:
        text = span.text
        if not text:
            continue
        if prev is not None:
            gap = span.bbox[0] - prev.bbox[2]
            prev_tail = prev.text[-1:] if prev.text else ""
            curr_head = text[:1]
            if (
                gap > max(2.0, span.size * 0.5)
                and prev_tail.isascii()
                and prev_tail.isalnum()
                and curr_head.isascii()
                and curr_head.isalnum()
            ):
                parts.append(" ")
        parts.append(text)
        prev = span
    return "".join(parts).strip()


def extract_pdf(data: bytes) -> ExtractedDocument:
    try:
        import fitz  # PyMuPDF
    except ImportError as exc:  # pragma: no cover - 依赖缺失时兜底
        raise ParseBusinessError("E_DEPENDENCY_MISSING", "缺少 PyMuPDF 依赖", 500) from exc

    try:
        doc = fitz.open(stream=data, filetype="pdf")
    except Exception as exc:
        raise ParseBusinessError("E_PARSE_FAILED", f"PDF 打开失败: {exc}", 422) from exc

    try:
        if getattr(doc, "needs_pass", False):
            raise ParseBusinessError("E_ENCRYPTED_PDF", "加密 PDF 无法解析", 422)

        spans: List[Span] = []
        text_lines: List[str] = []
        page_width = 0.0

        for page_index, page in enumerate(doc):
            try:
                page_dict = page.get_text("dict")
            except Exception:
                continue
            page_width = page_width or float(page.rect.width or 0.0)
            for block_index, block in enumerate(page_dict.get("blocks", [])):
                if block.get("type") != 0:
                    continue
                for line in block.get("lines", []):
                    line_spans: List[Span] = []
                    for raw in line.get("spans", []):
                        content = str(raw.get("text", ""))
                        if not content.strip():
                            continue
                        bbox_raw = raw.get("bbox") or (0.0, 0.0, 0.0, 0.0)
                        bbox = tuple(float(value) for value in bbox_raw)
                        span = Span(
                            text=content,
                            bbox=(bbox[0], bbox[1], bbox[2], bbox[3]),
                            size=float(raw.get("size", 0.0) or 0.0),
                            font=str(raw.get("font", "")),
                            bold=_is_bold(str(raw.get("font", "")), int(raw.get("flags", 0) or 0)),
                            page_no=page_index,
                            block_no=block_index,
                        )
                        spans.append(span)
                        line_spans.append(span)
                    if line_spans:
                        text_lines.append(_pdf_line_text(line_spans))

        text = "\n".join(text_lines)
        if len(text.strip()) < 50:
            raise ParseBusinessError(
                "E_SCANNED_PDF",
                "未提取到有效文本（疑似扫描件），本服务不支持 OCR",
                422,
            )
        return ExtractedDocument(
            spans=spans,
            tables=[],
            text=text,
            page_count=len(doc),
            doc_type="PDF",
            page_width=page_width,
        )
    finally:
        try:
            doc.close()
        except Exception:
            pass


def extract_docx(data: bytes) -> ExtractedDocument:
    try:
        from docx import Document  # python-docx
    except ImportError as exc:  # pragma: no cover
        raise ParseBusinessError("E_DEPENDENCY_MISSING", "缺少 python-docx 依赖", 500) from exc

    try:
        document = Document(io.BytesIO(data))
    except Exception as exc:
        raise ParseBusinessError("E_PARSE_FAILED", f"DOCX 打开失败: {exc}", 422) from exc

    spans: List[Span] = []
    text_lines: List[str] = []
    for index, paragraph in enumerate(document.paragraphs):
        content = (paragraph.text or "").strip()
        if not content:
            continue
        style_name = ""
        try:
            style_name = paragraph.style.name if paragraph.style is not None else ""
        except Exception:
            style_name = ""
        bold = any(bool(run.bold) for run in paragraph.runs)
        size = 0.0
        for run in paragraph.runs:
            if run.font is not None and run.font.size is not None:
                size = float(run.font.size.pt)
                break
        y0 = float(index * 20)
        spans.append(
            Span(
                text=content,
                bbox=(0.0, y0, 0.0, y0 + 15.0),
                size=size or 10.0,
                font="",
                bold=bold,
                page_no=0,
                block_no=index,
                style=style_name,
            )
        )
        text_lines.append(content)

    tables: List[List[List[str]]] = []
    for table in document.tables:
        rows: List[List[str]] = []
        for row in table.rows:
            cells: List[str] = []
            for cell in row.cells:
                parts = [(paragraph.text or "").strip() for paragraph in cell.paragraphs]
                cells.append(" ".join(part for part in parts if part))
            if any(cells):
                rows.append(cells)
        if rows:
            tables.append(rows)

    return ExtractedDocument(
        spans=spans,
        tables=tables,
        text="\n".join(text_lines),
        page_count=1,
        doc_type="DOCX",
        page_width=0.0,
    )


def extract_document(data: bytes, doc_type: str) -> ExtractedDocument:
    kind = (doc_type or "").strip().upper().lstrip(".")
    if kind == "PDF":
        return extract_pdf(data)
    if kind == "DOCX":
        return extract_docx(data)
    if kind == "DOC":
        raise ParseBusinessError("E_UNSUPPORTED_FORMAT", "暂不支持 .doc，请转换为 .docx", 415)
    raise ParseBusinessError("E_UNSUPPORTED_FORMAT", f"不支持的文件类型: {doc_type}", 415)
